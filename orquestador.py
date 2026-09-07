# -*- coding: utf-8 -*-
"""
CAPA 4 — Orquestador: recorre el Maestro de Clientes y procesa un cliente
(o todos) para un período dado, generando las 2 salidas en su carpeta.

No tiene nada de Streamlit — se puede probar por consola. `app.py` (la UI)
solo llama a las funciones de acá y muestra el resultado.

Nada de este módulo sabe reglas de negocio de ningún cliente: todo el
criterio de clasificación sale de los 4 archivos que cada cliente tiene en
su propia carpeta (ver conversor_extractos.py / importador_xubio.py).
"""
import io

import pdfplumber

import conversor_extractos as C
import drive_io
import importador_xubio as X

MESES_ES = ['ENERO', 'FEBRERO', 'MARZO', 'ABRIL', 'MAYO', 'JUNIO', 'JULIO',
            'AGOSTO', 'SEPTIEMBRE', 'OCTUBRE', 'NOVIEMBRE', 'DICIEMBRE']

# Nombres base (sin sufijo de cliente) de los 4 archivos de config -- ver
# contrato de datos en PLAN_CLAUDE_CODE_conciliacion.md.
_ARCHIVOS_CONFIG = dict(
    conceptos='CONCEPTO GASTO Y ASIGNACION CUENTA',
    proveedores='Listado y Ca Cte de Proveedores',
    vep='Vep Pagados',
    plan_cuentas='PLAN DE CUENTAS',
)


class ClienteSinConfig(Exception):
    """Al cliente le falta alguno de los 4 archivos de config en su carpeta."""


# ---------------------------------------------------------------------------
# Maestro de Clientes
# ---------------------------------------------------------------------------
def listar_clientes(maestro_id: str) -> list:
    """Lee la solapa 'Hoja 1' del Maestro de Clientes. Devuelve una lista de
    dicts: {razon_social, estado, link_carpeta, cuit}. Una fila sin
    RAZON SOCIAL se ignora; una fila sin LINK CARPETA se devuelve igual
    (con link_carpeta='') para que la UI la muestre como "sin carpeta
    cargada" en vez de desaparecer calladita."""
    archivo = drive_io.abrir_archivo_por_id(maestro_id)
    filas = drive_io.leer_sheet(archivo, 'Hoja 1')
    clientes = []
    for fila in filas:
        nombre = str(fila.get('RAZON SOCIAL') or '').strip()
        if not nombre:
            continue
        clientes.append(dict(
            razon_social=nombre,
            estado=str(fila.get('ESTADO') or '').strip(),
            link_carpeta=str(fila.get('LINK CARPETA') or '').strip(),
            cuit=str(fila.get('CUIT') or '').strip(),
        ))
    return clientes


def leer_bancos_maestro(maestro_id: str) -> dict:
    """Lee la solapa 'BANCOS' del Maestro (columnas repetidas -> se lee
    cruda, no como diccionario). Devuelve {razon_social -> [lista de
    códigos BANCO_... cargados, sin los huecos vacíos]}. Es solo
    informativo (qué monedas/bancos declaró tener el cliente); la cantidad
    real de cuentas sale de los PDFs que aparezcan + el CBU del Plan de
    Cuentas (ver importador_xubio.resolver_banco_por_cbu)."""
    archivo = drive_io.abrir_archivo_por_id(maestro_id)
    filas = drive_io.leer_sheet_crudo(archivo, 'BANCOS')
    out = {}
    for fila in filas[1:]:
        if not fila or not fila[0].strip():
            continue
        out[fila[0].strip()] = [c.strip() for c in fila[1:] if c.strip()]
    return out


# ---------------------------------------------------------------------------
# Config de UN cliente
# ---------------------------------------------------------------------------
def cargar_config_cliente(carpeta_id: str) -> dict:
    """Busca y lee los 4 archivos de config del cliente en su carpeta.
    Lanza ClienteSinConfig si falta alguno (sin esos 4 archivos no se puede
    clasificar nada con certeza)."""
    solapas_por_archivo = {}
    for clave, nombre_base in _ARCHIVOS_CONFIG.items():
        archivo = drive_io.buscar_archivo_por_nombre_base(carpeta_id, nombre_base)
        if archivo is None:
            raise ClienteSinConfig(f"No encontré el archivo '{nombre_base}' en la carpeta del cliente.")
        solapas_por_archivo[clave] = {
            s: drive_io.leer_sheet(archivo, s) for s in drive_io.listar_solapas(archivo)
        }

    return dict(
        conceptos=C.cargar_conceptos(solapas_por_archivo['conceptos']),
        empleados_socios=C.cargar_empleados_socios(solapas_por_archivo['proveedores']),
        proveedores=C.cargar_proveedores(solapas_por_archivo['proveedores']),
        vep=C.cargar_vep(solapas_por_archivo['vep']),
        plan_cuentas_c2=C.cargar_plan_cuentas(solapas_por_archivo['plan_cuentas']),
        plan_cuentas_x=X.cargar_plan_cuentas(solapas_por_archivo['plan_cuentas']),
    )


# ---------------------------------------------------------------------------
# Ubicar la subcarpeta del período dentro de la carpeta del cliente
# ---------------------------------------------------------------------------
def buscar_subcarpeta_periodo(carpeta_cliente_id: str, anio: int, mes: int):
    """Entra a <carpeta_cliente>/<AÑO>/<MES en español>. Devuelve el ID de
    Drive de esa subcarpeta, o None si no existe (nada cargado todavía para
    ese período — no es un error, el llamador lo maneja)."""
    sub_anio = drive_io.buscar_subcarpeta(carpeta_cliente_id, str(anio))
    if sub_anio is None:
        return None
    return drive_io.buscar_subcarpeta(sub_anio, MESES_ES[mes - 1])


# ---------------------------------------------------------------------------
# Procesar UN cliente para UN período
# ---------------------------------------------------------------------------
def procesar_cliente(nombre_cliente: str, link_carpeta: str, anio: int, mes: int,
                      forzar: bool = False) -> dict:
    """Procesa un cliente para un período: baja sus PDFs de ese mes, los
    clasifica con SU PROPIA config, arma los asientos de Xubio (todas las
    cuentas bancarias del cliente juntas, para poder resolver las
    transferencias entre cuentas propias) y sube los 2 archivos de salida a
    la carpeta del período.

    Nunca levanta excepción por datos faltantes: lo que falta queda anotado
    en el reporte devuelto, para no frenar el resto del lote por un
    cliente puntual. `forzar=False` (default) saltea el cliente si ya
    existe la salida de ese período (retomable)."""
    reporte = dict(cliente=nombre_cliente, ok=False, saltado=False, motivo='',
                    pdfs_procesados=[], pdfs_con_error=[], controles=[], cobertura=[])
    try:
        carpeta_id = drive_io.abrir_carpeta_cliente(link_carpeta)
    except Exception as e:
        reporte['motivo'] = f"No pude abrir la carpeta del cliente: {e}"
        return reporte

    sub_periodo = buscar_subcarpeta_periodo(carpeta_id, anio, mes)
    if sub_periodo is None:
        reporte['motivo'] = (f"No existe la subcarpeta {anio}/{MESES_ES[mes - 1].title()} "
                              f"en la carpeta del cliente todavía.")
        return reporte

    nombre_importador = f"IMPORTADOR ASIENTOS - {nombre_cliente} - {mes:02d}.{anio}.xlsx"
    nombre_papel = f"Papel de trabajo - {nombre_cliente} - {mes:02d}.{anio}.xlsx"

    if not forzar and drive_io.existe_archivo(sub_periodo, nombre_importador):
        reporte['ok'] = True
        reporte['saltado'] = True
        reporte['motivo'] = "Ya estaba procesado para este período (se saltea)."
        return reporte

    pdfs = drive_io.listar_pdfs(sub_periodo)
    if not pdfs:
        reporte['motivo'] = f"No hay ningún PDF cargado en {anio}/{MESES_ES[mes - 1].title()}."
        return reporte

    try:
        config = cargar_config_cliente(carpeta_id)
    except ClienteSinConfig as e:
        reporte['motivo'] = str(e)
        return reporte

    resultados = []
    for pdf_info in pdfs:
        try:
            contenido = drive_io.descargar_pdf(pdf_info)
            with pdfplumber.open(io.BytesIO(contenido)) as pdf:
                _, lector = C.detectar_banco(pdf)
                salida = lector(pdf)
            for res in (salida if isinstance(salida, list) else [salida]):
                C.clasificar(res, config['conceptos'], config['empleados_socios'],
                             config['proveedores'], config['vep'], config['plan_cuentas_c2'])
                resultados.append(res)
            reporte['pdfs_procesados'].append(pdf_info['name'])
        except Exception as e:
            reporte['pdfs_con_error'].append((pdf_info['name'], str(e)))

    if not resultados:
        reporte['motivo'] = "Ningún PDF de este período se pudo procesar (ver pdfs_con_error)."
        return reporte

    reporte['controles'] = [dict(cuenta=res.cuenta, **C.verificar_control(res)) for res in resultados]
    reporte['cobertura'] = X.control_cobertura(resultados, config['plan_cuentas_x'])

    asientos, revisar, internas = X.construir_asientos(resultados, config['plan_cuentas_x'])
    reporte['pendientes'] = len(revisar)

    buf_importador = io.BytesIO()
    X.generar_excel_importador(asientos, buf_importador)

    # Papel de trabajo: resumen de asientos + pendientes + control, MÁS una
    # hoja de detalle movimiento por movimiento por cada cuenta bancaria
    # (el "extracto convertido" completo, para poder revisar cada línea --
    # no solo las que quedaron pendientes).
    wb_papel = X.construir_wb_papel_trabajo(resultados, config['plan_cuentas_x'], asientos, revisar, internas)
    usados = set()
    hoja_por_res = {}
    for res in resultados:
        ws = wb_papel.create_sheet(C._sanitizar_hoja(res.cuenta, usados))
        C._escribir_hoja(ws, res)
        hoja_por_res[id(res)] = ws.title

    # Link "ver detalle": desde cada fila de A REVISAR a la fila exacta del
    # movimiento en la hoja de detalle de su cuenta (_escribir_hoja pone el
    # primer movimiento en la fila 4, uno por fila, en el mismo orden que
    # res.movimientos -- por eso se puede calcular la fila sin repetir nada).
    if 'A REVISAR' in wb_papel.sheetnames:
        from openpyxl.styles import Font
        from openpyxl.worksheet.hyperlink import Hyperlink
        wa = wb_papel['A REVISAR']
        wa.cell(1, 8, 'VER DETALLE').font = Font(bold=True)
        for fila_a_revisar, (res_pend, mov_pend, _motivo) in enumerate(revisar, start=2):
            nombre_hoja = hoja_por_res.get(id(res_pend))
            indice = next((i for i, m in enumerate(res_pend.movimientos) if m is mov_pend), None)
            if nombre_hoja is None or indice is None:
                continue
            fila_detalle = indice + 4
            celda = wa.cell(fila_a_revisar, 8, 'Ver fila →')
            # location= (no target=) es el link INTERNO "de verdad" -- sin
            # relación externa. La forma target="#'Hoja'!A1" queda marcada
            # como link externo en el xlsx: Excel la tolera, pero Google
            # Sheets no la sigue.
            celda.hyperlink = Hyperlink(ref=celda.coordinate, location=f"'{nombre_hoja}'!A{fila_detalle}")
            celda.font = Font(color='0563C1', underline='single')
        wa.column_dimensions['H'].width = 16
    buf_papel = io.BytesIO()
    wb_papel.save(buf_papel)

    drive_io.escribir_excel_en_carpeta(sub_periodo, nombre_importador, buf_importador.getvalue())
    drive_io.escribir_excel_en_carpeta(sub_periodo, nombre_papel, buf_papel.getvalue())

    reporte['ok'] = True
    return reporte
